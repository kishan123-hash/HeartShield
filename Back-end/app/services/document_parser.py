import io
import os
import re
import uuid
import logging
from pathlib import Path
from typing import Tuple, Optional, Dict, Any

from pypdf import PdfReader
from docx import Document

from app.services.gemini_service import gemini_service

logger = logging.getLogger("document_parser")

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB


class DocumentParserService:
    """
    Secure medical document validation and text extraction:
    - Content-based file signature validation
    - Multi-format extraction (PDF, DOCX, Images)
    - Automated clinical summary generation
    """

    ALLOWED_MIME_TYPES = {
        "application/pdf": [b"%PDF-"],
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": [b"PK\x03\x04"],
        "image/png": [b"\x89PNG\r\n\x1a\n"],
        "image/jpeg": [b"\xff\xd8\xff"],
        "image/webp": [b"RIFF"]
    }

    @classmethod
    def validate_file(cls, file_bytes: bytes, filename: str, mime_type: str) -> Tuple[bool, str]:
        """
        Validates file size, extension, and magic byte signature.
        """
        if len(file_bytes) == 0:
            return False, "Uploaded file is empty."

        if len(file_bytes) > MAX_FILE_SIZE_BYTES:
            return False, f"File size exceeds the 10 MB limit ({len(file_bytes)/(1024*1024):.1f} MB)."

        ext = Path(filename).suffix.lower()
        valid_exts = [".pdf", ".docx", ".png", ".jpg", ".jpeg", ".webp"]
        if ext not in valid_exts:
            return False, f"Unsupported file extension '{ext}'. Allowed: PDF, DOCX, PNG, JPG, WEBP."

        # Verify signature
        matched_sig = False
        for expected_mime, signatures in cls.ALLOWED_MIME_TYPES.items():
            for sig in signatures:
                if file_bytes.startswith(sig):
                    matched_sig = True
                    break

        if not matched_sig:
            return False, "File content signature does not match allowed document or image types."

        return True, "Valid"

    @classmethod
    def extract_text(cls, file_bytes: bytes, filename: str, mime_type: str) -> str:
        """
        Extracts raw textual and tabular content from PDF, DOCX, or Image.
        """
        ext = Path(filename).suffix.lower()

        if ext == ".pdf":
            try:
                reader = PdfReader(io.BytesIO(file_bytes))
                extracted = []
                for i, page in enumerate(reader.pages):
                    page_text = page.extract_text()
                    if page_text:
                        extracted.append(f"--- Page {i+1} ---\n{page_text}")
                
                full_text = "\n\n".join(extracted).strip()
                if not full_text:
                    return "[Scanned PDF detected without embedded text. Visual review recommended.]"
                return full_text
            except Exception as e:
                logger.error(f"PDF extraction error: {e}")
                return f"[Error parsing PDF document: {e}]"

        elif ext == ".docx":
            try:
                doc = Document(io.BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                # Extract tables
                table_texts = []
                for table in doc.tables:
                    for row in table.rows:
                        row_vals = [cell.text.strip() for cell in row.cells]
                        table_texts.append(" | ".join(row_vals))
                
                full_text = "\n".join(paragraphs)
                if table_texts:
                    full_text += "\n\n--- Tables ---\n" + "\n".join(table_texts)
                return full_text.strip()
            except Exception as e:
                logger.error(f"DOCX extraction error: {e}")
                return f"[Error parsing DOCX document: {e}]"

        elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
            # Image OCR via Gemini Vision
            return gemini_service.extract_report_with_vision(file_bytes, mime_type)

        return ""

    @classmethod
    def generate_clinical_summary(cls, extracted_text: str) -> str:
        """
        Generates an educational summary of the extracted laboratory results.
        """
        if not extracted_text or len(extracted_text.strip()) < 20:
            return "Document processed. No specific clinical laboratory metrics detected."

        prompt = (
            "Summarize the key findings from this patient medical report in 2 clear, accessible paragraphs. "
            "Highlight any mentioned cardiovascular markers (e.g., blood pressure, lipid panel, blood glucose) "
            "and state clearly whether values are within standard reference ranges. "
            "Do not prescribe medications or offer definitive diagnoses.\n\n"
            f"REPORT CONTENT:\n{extracted_text[:3500]}"
        )

        try:
            if gemini_service.client:
                response = gemini_service.client.models.generate_content(
                    model=gemini_service.model_name,
                    contents=prompt
                )
                return response.text or "Clinical summary generated."
        except Exception as e:
            logger.error(f"Summary generation error: {e}")

        # Fallback simple summary
        lines = [line.strip() for line in extracted_text.splitlines() if line.strip()]
        return f"Medical report uploaded successfully ({len(lines)} lines of clinical data recorded)."


document_parser = DocumentParserService()
